import os, re
from pptx import Presentation
from tqdm import tqdm
from config import PPT_DIR, PPT_TXT_DIR
from utils import logger, check_dir

def clean_ppt_text(text):
    text = re.sub(r"第\s*\d+\s*页|页码\s*\d+", "", text)
    text = re.sub(r"\n+", "\n", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def extract_slide_text(slide):
    parts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
    return "\n".join(parts)

def should_parse(src_path, dst_path):
    if not os.path.exists(dst_path):
        return True
    return os.path.getmtime(src_path) > os.path.getmtime(dst_path)

def parse_all_ppt(force=False):
    check_dir(PPT_TXT_DIR)
    if not os.path.exists(PPT_DIR):
        logger.error(f"PPT文件夹不存在: {PPT_DIR}")
        return

    files = [f for f in os.listdir(PPT_DIR) if f.endswith(".pptx")]
    if not files:
        logger.warning("未找到PPT文件")
        return

    for name in tqdm(files, desc="解析PPT"):
        src_path = os.path.join(PPT_DIR, name)
        out_name = os.path.splitext(name)[0] + ".txt"
        dst_path = os.path.join(PPT_TXT_DIR, out_name)

        if not force and not should_parse(src_path, dst_path):
            logger.info(f"⏭️ 跳过 {name}，TXT 文件已是最新")
            continue

        try:
            prs = Presentation(src_path)
            pages = []
            for i, slide in enumerate(prs.slides, 1):
                txt = extract_slide_text(slide)
                if txt:
                    pages.append(f"【PPT第{i}页】\n{txt}")
            raw = "\n\n".join(pages)
            clean = clean_ppt_text(raw)
            with open(dst_path, "w", encoding="utf-8") as f:
                f.write(clean)
            logger.info(f"✅ 完成: {name}")
        except Exception as e:
            logger.error(f"❌ 失败 {name}: {e}")

    logger.info("✅ PPT解析完成")

if __name__ == "__main__":
    parse_all_ppt()
