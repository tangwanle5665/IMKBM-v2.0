import os, re
from tqdm import tqdm
from pymupdf import open as fitz_open
from docx import Document
from config import PDF_DIR, TXT_DIR
from utils import logger, check_dir

def clean_text(text: str) -> str:
    text = re.sub(r"^\s*\d+\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"第\s*\d+\s*页(\s*/\s*共\s*\d+\s*页)?", "", text)
    text = re.sub(r"\n\s*\n", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()

def parse_pdf_to_text(pdf_path):
    doc = fitz_open(pdf_path)
    parts = [page.get_text() for page in doc if page.get_text()]
    doc.close()
    return "\n".join(parts)

def parse_docx_to_text(docx_path):
    doc = Document(docx_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt: parts.append(txt)
    return "\n".join(parts)

def should_parse(src_path, dst_path):
    """判断是否需要解析：目标文件不存在，或源文件比目标文件新"""
    if not os.path.exists(dst_path):
        return True
    src_mtime = os.path.getmtime(src_path)
    dst_mtime = os.path.getmtime(dst_path)
    return src_mtime > dst_mtime

def parse_files_to_txt(force=False):
    """
    force=True 时强制重新解析所有文件；
    force=False（默认）时只解析新增或更新过的文件。
    """
    check_dir(TXT_DIR)
    if not os.path.exists(PDF_DIR):
        logger.error(f"教材文件夹不存在: {PDF_DIR}")
        return

    files = os.listdir(PDF_DIR)
    pdfs = [f for f in files if f.lower().endswith(".pdf")]
    docxs = [f for f in files if f.lower().endswith(".docx")]

    # 解析 PDF
    for name in tqdm(pdfs, desc="解析PDF"):
        src_path = os.path.join(PDF_DIR, name)
        out_name = os.path.splitext(name)[0] + ".txt"
        dst_path = os.path.join(TXT_DIR, out_name)

        if not force and not should_parse(src_path, dst_path):
            logger.info(f"⏭️ 跳过 {name}，TXT 文件已是最新")
            continue

        try:
            raw = parse_pdf_to_text(src_path)
            clean = clean_text(raw)
            with open(dst_path, "w", encoding="utf-8") as f:
                f.write(clean)
            logger.info(f"✅ PDF完成: {name}")
        except Exception as e:
            logger.error(f"❌ PDF失败 {name}: {e}")

    # 解析 Word
    for name in tqdm(docxs, desc="解析Word"):
        src_path = os.path.join(PDF_DIR, name)
        out_name = os.path.splitext(name)[0] + ".txt"
        dst_path = os.path.join(TXT_DIR, out_name)

        if not force and not should_parse(src_path, dst_path):
            logger.info(f"⏭️ 跳过 {name}，TXT 文件已是最新")
            continue

        try:
            raw = parse_docx_to_text(src_path)
            clean = clean_text(raw)
            with open(dst_path, "w", encoding="utf-8") as f:
                f.write(clean)
            logger.info(f"✅ Word完成: {name}")
        except Exception as e:
            logger.error(f"❌ Word失败 {name}: {e}")

    logger.info("✅ 全部教材解析完成")

if __name__ == "__main__":
    parse_files_to_txt()
